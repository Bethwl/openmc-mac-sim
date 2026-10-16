"""
sim_config.py -- Doc cau hinh mo phong TU 1 FILE EXCEL DUY NHAT (Config_Simulation.xlsx)
Dung CHUNG boi run_all.py, Run_nosample.py, Run_sample.py cho Paper 10 (Paper 10.1 - PNE 2022).

Thuy tinh Gadolinium Silicoborate gom 8 mau (S1 -> S8):
  - S1: Li-B-O-Al-Si (khong Gd)
  - S2..S4: Li-B-O-Al-Si-Gd (5, 10, 15 mol% Gd2O3)
  - S5..S8: B-O-Si-Ca-Gd (10 mol% CaO, 20..35 mol% Gd2O3)
He nguyen to day du: Li, B, O, Al, Si, Ca, Gd.

Sheet Excel:
  - "Sample_Composition"  : Sample | Density_g_cm3 | Li_wo | B_wo | O_wo | Al_wo | Si_wo | Ca_wo | Gd_wo
  - "Energy_Thickness"    : Energy(keV) | S1..S8 Be_day_cm (23 muc nang luong tu 15 keV den 15 MeV)
"""
import os
import sys

# 8 mau thuy tinh chuan cua Paper 10
SAMPLES = ['S1', 'S2', 'S3', 'S4', 'S5', 'S6', 'S7', 'S8']

# He nguyen to day du cho ca 8 mau
REQUIRED_SAMPLE_COLUMNS = ['Sample', 'Density_g_cm3', 'Li_wo', 'B_wo', 'O_wo', 'Al_wo', 'Si_wo', 'Ca_wo', 'Gd_wo']
ALL_ELEMENT_KEYS = ['Li', 'B', 'O', 'Al', 'Si', 'Ca', 'Gd']


def _fail(msg):
    print(f"  [LOI] {msg}")
    sys.exit(1)


def _load_sheet_rows(path, sheet_name):
    """Doc toan bo cac dong trong 1 sheet. Bao loi va dung neu thieu file/sheet."""
    if not path:
        _fail("Khong truyen file config Excel (CONFIG_EXCEL rong/None).")
    if not os.path.isfile(path):
        _fail(f"Khong tim thay file config Excel: {path}")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        _fail(f"Khong mo duoc file Excel {path}: {e}")
    if sheet_name not in wb.sheetnames:
        _fail(f"File {path} thieu sheet '{sheet_name}' (cac sheet hien co: {wb.sheetnames}).")
    ws = wb[sheet_name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        _fail(f"Sheet '{sheet_name}' trong {path} rong.")
    return rows


def load_sample_composition(path):
    """Doc sheet Sample_Composition -- 8 mau S1..S8, he nguyen to Li/B/O/Al/Si/Ca/Gd."""
    rows = _load_sheet_rows(path, 'Sample_Composition')
    header = [str(h).strip() if h else '' for h in rows[0]]
    try:
        idx = {name: header.index(name) for name in REQUIRED_SAMPLE_COLUMNS}
    except ValueError as e:
        _fail(f"Sheet Sample_Composition trong {path} thieu cot: {e}")

    result = {}
    for row in rows[1:]:
        if not row or row[idx['Sample']] is None:
            continue
        sname = str(row[idx['Sample']]).strip().upper()
        if sname not in SAMPLES:
            continue
        try:
            entry = {
                'density': float(row[idx['Density_g_cm3']]),
                'Li': float(row[idx['Li_wo']]),
                'B':  float(row[idx['B_wo']]),
                'O':  float(row[idx['O_wo']]),
                'Al': float(row[idx['Al_wo']]),
                'Si': float(row[idx['Si_wo']]),
                'Ca': float(row[idx['Ca_wo']]),
                'Gd': float(row[idx['Gd_wo']]),
            }
        except (ValueError, TypeError) as e:
            _fail(f"Sample_Composition dong '{sname}' trong {path}: gia tri khong hop le ({e}).")

        # Kiem tong ty le khoi luong phai ~ 1.0
        total_wo = sum(entry[k] for k in ALL_ELEMENT_KEYS)
        if abs(total_wo - 1.0) > 0.01:
            _fail(f"Sample_Composition dong '{sname}' trong {path}: tong ty le khoi luong = {total_wo:.5f}, phai ~1.0.")

        result[sname] = entry

    missing_samples = [s for s in SAMPLES if s not in result]
    if missing_samples:
        _fail(f"Sheet Sample_Composition trong {path} thieu du lieu cho mau: {missing_samples}")

    print(f"  [OK] Da nap du Sample_Composition ({len(result)}/{len(SAMPLES)} mau) tu {path}")
    return result


def load_energy_thickness(path):
    """Doc sheet 'Energy_Thickness' dang Energy(keV)|S1|...|S8.
    Tra ve: {E_keV: [ {S1: t1, ..., S8: t8}, ... ] }
    Moi nang luong co the co nhieu bo be day (dong 1, dong 2, ...).
    """
    rows = _load_sheet_rows(path, 'Energy_Thickness')
    if len(rows) < 3:
        _fail(f"Sheet Energy_Thickness trong {path} qua it dong (can >= 3 dong: 2 dong header + du lieu).")

    result = {}
    n_rows_loaded = 0
    for row_idx, row in enumerate(rows[2:], start=3):
        if not row or row[0] is None or str(row[0]).strip() == '':
            continue
        try:
            e_kev = float(row[0])
        except (ValueError, TypeError):
            _fail(f"Dong {row_idx} trong {path}: Energy khong hop le ({row[0]!r}).")

        vals = {}
        for i, s in enumerate(SAMPLES):
            cell = row[1 + i] if len(row) > 1 + i else None
            if cell is None or str(cell).strip() == '':
                _fail(f"Dong {row_idx} (E={e_kev} keV) trong {path}: thieu be day mau {s}.")
            try:
                t_val = float(cell)
            except (ValueError, TypeError):
                _fail(f"Dong {row_idx} (E={e_kev} keV) trong {path}: be day mau {s} khong phai so ({cell!r}).")
            if t_val <= 0:
                _fail(f"Dong {row_idx} (E={e_kev} keV) trong {path}: be day mau {s} phai > 0 (dang la {t_val}).")
            vals[s] = t_val

        result.setdefault(e_kev, []).append(vals)
        n_rows_loaded += 1

    if not result:
        _fail(f"Sheet Energy_Thickness trong {path} khong co dong du lieu nao hop le.")

    n_dup = sum(1 for v in result.values() if len(v) > 1)
    print(f"  [OK] Da nap du {n_rows_loaded} dong du lieu ({len(result)} muc nang luong, {n_dup} muc co nhieu be day) tu Energy_Thickness ({path})")
    return result
